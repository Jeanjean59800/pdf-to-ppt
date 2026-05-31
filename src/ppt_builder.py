import hashlib
import json
import os
import re
import time
import urllib.request
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt

from src.document_analysis import parse_slide_plan
from src.types import PipelineData, VisualAsset


SLIDE_LAYOUTS = {
    "text_only",
    "text_image_right",
    "image_focus",
    "two_images_compare",
    "chart_focus",
    "section_break",
}


def load_presentation_template(output_path: Path) -> Presentation:
    template_candidates = []

    configured_template = os.getenv("PPTX_TEMPLATE_PATH")
    if configured_template:
        template_candidates.append(Path(configured_template))

    project_dir = Path(__file__).resolve().parents[1]
    template_candidates.extend(
        [
            project_dir / "template.pptx",
            project_dir / "templates" / "template.pptx",
            project_dir / "templates" / "presentation_template.pptx",
            output_path.with_name("template.pptx"),
        ]
    )

    seen_paths: set[Path] = set()
    for candidate in template_candidates:
        resolved_candidate = candidate.expanduser()
        if not resolved_candidate.is_absolute():
            resolved_candidate = (project_dir / resolved_candidate).resolve()
        else:
            resolved_candidate = resolved_candidate.resolve()

        if resolved_candidate in seen_paths:
            continue
        seen_paths.add(resolved_candidate)

        if resolved_candidate.is_file():
            return Presentation(str(resolved_candidate))

    return Presentation()


def build_editable_pptx(
    data: PipelineData,
    slide_plan_text: str,
    output_path: Path,
    max_slides: int,
    enable_fal_image_generation: bool = True,
) -> Path:
    thesis, slides = parse_slide_plan(slide_plan_text)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = load_presentation_template(output_path)
    add_title_slide(prs, data.pdf_path.stem.replace("_", " ").title(), "Generated editable deck skeleton")
    add_bullet_slide(prs, "Document Thesis", [thesis or data.thesis_summary])
    add_inventory_slide(prs, data)

    available_visuals = list(data.cropped_visuals)
    available_plots = list(data.generated_plot_paths)
    generated_visuals_dir = output_path.parent / "generated_visuals"
    for index, slide in enumerate(slides[:max_slides], start=1):
        render_slide_from_plan(
            prs=prs,
            slide_number=index,
            slide=slide,
            visuals=available_visuals,
            plots=available_plots,
            generated_visuals_dir=generated_visuals_dir,
            enable_fal_image_generation=enable_fal_image_generation,
        )

    prs.save(output_path)
    return output_path


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for bullet in bullets:
        paragraph = body.add_paragraph()
        paragraph.text = bullet
        paragraph.font.size = Pt(18)


def add_inventory_slide(prs: Presentation, data: PipelineData) -> None:
    add_bullet_slide(
        prs,
        "Available Data",
        [
            f"PDF: {data.pdf_path.name}",
            f"Raw text: {data.raw_text_path}",
            f"Rendered page images: {len(data.page_images)}",
            f"Candidate cropped visuals: {len(data.cropped_visuals)}",
            f"Generated plots: {len(data.generated_plot_paths)}",
            "Manual refinement: validate slide logic, figure relevance, and wording.",
        ],
    )


def render_slide_from_plan(
    prs: Presentation,
    slide_number: int,
    slide: dict,
    visuals: list[VisualAsset],
    plots: list[Path],
    generated_visuals_dir: Path,
    enable_fal_image_generation: bool,
) -> None:
    title = f"{slide_number}. {slide.get('title', 'Untitled')}"[:80]
    content_md = str(slide.get("content_md", "")).strip()
    visual_note = str(slide.get("visual", "text only"))
    layout = normalize_layout(slide.get("layout", ""))
    primary_asset, secondary_asset = choose_assets_for_slide(
        slide,
        visuals,
        plots,
        generated_visuals_dir,
        enable_fal_image_generation,
    )

    if layout == "section_break":
        add_section_break_slide(prs, title, first_content_line(content_md))
        return
    if layout == "image_focus":
        add_image_focus_slide(prs, title, content_md, primary_asset, visual_note)
        return
    if layout == "two_images_compare":
        add_two_images_compare_slide(prs, title, content_md, primary_asset, secondary_asset, visual_note)
        return
    if layout == "chart_focus":
        add_chart_focus_slide(prs, title, content_md, primary_asset, visual_note)
        return
    if layout == "text_image_right":
        add_text_image_right_slide(prs, title, content_md, primary_asset, visual_note)
        return
    add_text_only_slide(prs, title, content_md, visual_note)


def normalize_layout(layout: str) -> str:
    normalized = str(layout).strip().lower().replace("-", "_").replace(" ", "_")
    return normalized if normalized in SLIDE_LAYOUTS else "text_only"


def choose_assets_for_slide(
    slide: dict,
    visuals: list[VisualAsset],
    plots: list[Path],
    generated_visuals_dir: Path,
    enable_fal_image_generation: bool,
) -> tuple[VisualAsset | None, VisualAsset | None]:
    requested_layout = normalize_layout(slide.get("layout", ""))
    visual_request = str(slide.get("visual", "")).lower()

    if requested_layout == "chart_focus" and plots:
        primary_plot = choose_plot_for_slide(slide, plots)
        return path_to_visual_asset(primary_plot, "generated plot"), None

    if requested_layout == "two_images_compare":
        assets = choose_multiple_visuals_for_slide(
            slide,
            visuals,
            plots,
            generated_visuals_dir,
            count=2,
            enable_fal_image_generation=enable_fal_image_generation,
        )
        if len(assets) == 2:
            return assets[0], assets[1]
        if len(assets) == 1:
            return assets[0], None
        return None, None

    if "extracted" in visual_request and visuals:
        matched_visual = choose_semantic_visual_match(slide, visuals)
        if matched_visual is not None:
            return matched_visual, None

    if requested_layout == "chart_focus" and not plots and visuals:
        matched_visual = choose_semantic_visual_match(slide, visuals)
        return matched_visual, None

    if enable_fal_image_generation and any(
        hint in visual_request for hint in ("generated illustration", "simple diagram", "extracted")
    ):
        return generate_supporting_illustration(slide, generated_visuals_dir), None
    return None, None


def choose_multiple_visuals_for_slide(
    slide: dict,
    visuals: list[VisualAsset],
    plots: list[Path],
    generated_visuals_dir: Path,
    count: int,
    enable_fal_image_generation: bool,
) -> list[VisualAsset]:
    selected: list[VisualAsset] = []

    while plots and len(selected) < count and "chart" in str(slide.get("visual", "")).lower():
        selected.append(path_to_visual_asset(plots.pop(0), "generated plot"))

    while visuals and len(selected) < count:
        match = choose_semantic_visual_match(slide, visuals)
        if match is None:
            break
        selected.append(match)

    while enable_fal_image_generation and len(selected) < count:
        generated = generate_supporting_illustration(slide, generated_visuals_dir)
        if generated is None:
            break
        selected.append(generated)

    return selected


def choose_plot_for_slide(slide: dict, plots: list[Path]) -> Path:
    slide_tokens = build_slide_visual_tokens(slide)
    if not slide_tokens:
        return plots.pop(0)

    best_index = 0
    best_score = -1.0
    for index, plot_path in enumerate(plots):
        score = len(slide_tokens & tokenize_for_matching(plot_path.stem.replace("_", " ")))
        if score > best_score:
            best_score = score
            best_index = index
    return plots.pop(best_index)


def path_to_visual_asset(path: Path, source: str) -> VisualAsset:
    return VisualAsset(path=path, source=source, description=path.stem.replace("_", " "))


def add_text_only_slide(prs: Presentation, title: str, content_md: str, visual_note: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    write_markdown(body, content_md, visual_note)


def add_text_image_right_slide(
    prs: Presentation,
    title: str,
    content_md: str,
    asset: VisualAsset | None,
    visual_note: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_textbox(slide, title)
    add_markdown_textbox(slide, Inches(0.6), Inches(1.5), Inches(5.2), Inches(5.2), content_md, visual_note)
    add_media_or_placeholder(slide, asset, Inches(6.0), Inches(1.5), Inches(3.2), Inches(4.8), visual_note)


def add_image_focus_slide(
    prs: Presentation,
    title: str,
    content_md: str,
    asset: VisualAsset | None,
    visual_note: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_textbox(slide, title)
    add_media_or_placeholder(slide, asset, Inches(0.7), Inches(1.4), Inches(8.6), Inches(4.4), visual_note)
    add_markdown_textbox(slide, Inches(0.7), Inches(5.95), Inches(8.6), Inches(1.0), content_md, visual_note=None)


def add_two_images_compare_slide(
    prs: Presentation,
    title: str,
    content_md: str,
    primary_asset: VisualAsset | None,
    secondary_asset: VisualAsset | None,
    visual_note: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_textbox(slide, title)
    add_media_or_placeholder(slide, primary_asset, Inches(0.6), Inches(1.5), Inches(4.35), Inches(3.6), "Left comparison slot")
    add_media_or_placeholder(
        slide,
        secondary_asset,
        Inches(5.05),
        Inches(1.5),
        Inches(4.35),
        Inches(3.6),
        "Right comparison slot",
    )
    add_markdown_textbox(slide, Inches(0.6), Inches(5.3), Inches(8.8), Inches(1.4), content_md, visual_note)


def add_chart_focus_slide(
    prs: Presentation,
    title: str,
    content_md: str,
    asset: VisualAsset | None,
    visual_note: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_textbox(slide, title)
    add_media_or_placeholder(slide, asset, Inches(0.7), Inches(1.5), Inches(5.7), Inches(4.7), visual_note)
    add_markdown_textbox(slide, Inches(6.55), Inches(1.6), Inches(2.7), Inches(4.6), content_md, None)


def add_section_break_slide(prs: Presentation, title: str, subtitle: str | None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(1.35), prs.slide_width, Inches(3.2))
    banner.fill.solid()
    banner.fill.fore_color.rgb = _rgb(31, 56, 100)
    banner.line.color.rgb = _rgb(31, 56, 100)
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(8.0), Inches(1.0))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(28)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = _rgb(255, 255, 255)
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.1), Inches(8.2), Inches(0.8))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_frame.paragraphs[0].font.size = Pt(16)
        subtitle_frame.paragraphs[0].font.color.rgb = _rgb(255, 255, 255)


def add_title_textbox(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8.8), Inches(0.75))
    frame = box.text_frame
    frame.text = title
    frame.paragraphs[0].font.size = Pt(24)
    frame.paragraphs[0].font.bold = True


def add_markdown_textbox(slide, left, top, width, height, content_md: str, visual_note: str | None) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    write_markdown(box.text_frame, content_md, visual_note)


def write_markdown(text_frame, content_md: str, visual_note: str | None) -> None:
    text_frame.clear()
    blocks = parse_markdown_blocks(content_md)
    if not blocks:
        blocks = [{"type": "body", "text": "Content to refine manually."}]
    for block in blocks:
        paragraph = text_frame.add_paragraph()
        paragraph.text = block["text"]
        if block["type"] == "heading":
            paragraph.font.size = Pt(18)
            paragraph.font.bold = True
        elif block["type"] == "bullet":
            paragraph.font.size = Pt(16)
            paragraph.level = 0
        else:
            paragraph.font.size = Pt(15)
    if visual_note:
        note = text_frame.add_paragraph()
        note.text = f"Visual direction: {visual_note}"
        note.font.size = Pt(11)


def parse_markdown_blocks(content_md: str) -> list[dict[str, str]]:
    blocks = []
    for raw_line in content_md.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            blocks.append({"type": "heading", "text": clean_inline_markdown(line[3:].strip())})
            continue
        if line.startswith("- "):
            blocks.append({"type": "bullet", "text": clean_inline_markdown(line[2:].strip())})
            continue
        blocks.append({"type": "body", "text": clean_inline_markdown(line)})
    return blocks


def clean_inline_markdown(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("*", "")
    return text.strip()


def first_content_line(content_md: str) -> str | None:
    for block in parse_markdown_blocks(content_md):
        if block["text"]:
            return block["text"]
    return None


def add_media_or_placeholder(slide, asset: VisualAsset | None, left, top, width, height, note: str) -> None:
    if asset:
        try:
            slide.shapes.add_picture(str(asset.path), left, top, width=width, height=height)
            return
        except Exception:
            pass

    placeholder = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    placeholder.fill.solid()
    placeholder.fill.fore_color.rgb = _rgb(241, 244, 248)
    placeholder.line.color.rgb = _rgb(190, 198, 210)
    frame = placeholder.text_frame
    frame.text = note or "Media slot"
    frame.paragraphs[0].font.size = Pt(14)


def _rgb(red: int, green: int, blue: int):
    from pptx.dml.color import RGBColor

    return RGBColor(red, green, blue)


def choose_semantic_visual_match(slide: dict, visuals: list[VisualAsset]) -> VisualAsset | None:
    slide_tokens = build_slide_visual_tokens(slide)
    if not slide_tokens:
        return visuals.pop(0) if visuals else None

    best_index = -1
    best_score = 0.0
    for index, visual in enumerate(visuals):
        score = score_visual_match(slide_tokens, visual)
        if score > best_score:
            best_score = score
            best_index = index

    if best_index >= 0:
        return visuals.pop(best_index)
    return visuals.pop(0) if visuals else None


def build_slide_visual_tokens(slide: dict) -> set[str]:
    parts = [
        str(slide.get("title", "")),
        str(slide.get("layout", "")),
        str(slide.get("visual", "")),
        str(slide.get("content_md", "")),
    ]
    return tokenize_for_matching(" ".join(parts))


def score_visual_match(slide_tokens: set[str], visual: VisualAsset) -> float:
    description_tokens = tokenize_for_matching(visual.description)
    source_tokens = tokenize_for_matching(visual.source)
    name_tokens = tokenize_for_matching(visual.path.stem.replace("_", " "))

    score = 0.0
    if description_tokens:
        score += len(slide_tokens & description_tokens) * 3.0
    if source_tokens:
        score += len(slide_tokens & source_tokens) * 1.5
    if name_tokens:
        score += len(slide_tokens & name_tokens)
    return score


def tokenize_for_matching(text: str) -> set[str]:
    stopwords = {
        "a",
        "an",
        "and",
        "as",
        "by",
        "for",
        "from",
        "in",
        "of",
        "on",
        "or",
        "the",
        "to",
        "with",
        "visual",
        "figure",
        "image",
        "slide",
        "candidate",
        "generated",
        "illustration",
        "diagram",
        "page",
        "text",
        "only",
        "layout",
    }
    tokens = re.findall(r"[a-z0-9]+", text.lower())
    return {token for token in tokens if len(token) > 2 and token not in stopwords}


def generate_supporting_illustration(slide: dict, output_dir: Path) -> VisualAsset | None:
    fal_key = os.getenv("FAL_KEY")
    if not fal_key:
        return None

    prompt = build_illustration_prompt(slide)
    prompt_hash = hashlib.sha1(prompt.encode("utf-8")).hexdigest()[:12]
    image_path = output_dir / f"{slugify(slide.get('title', 'slide'))}-{prompt_hash}.png"
    if image_path.exists():
        return VisualAsset(path=image_path, source="fal", description=prompt)

    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        request = urllib.request.Request(
            "https://queue.fal.run/fal-ai/flux/schnell",
            data=json.dumps(
                {
                    "prompt": prompt,
                    "image_size": "landscape_4_3",
                    "num_images": 1,
                    "num_inference_steps": 4,
                }
            ).encode("utf-8"),
            headers={
                "Authorization": f"Key {fal_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=20) as response:
            payload = json.loads(response.read().decode("utf-8"))

        response_url = payload.get("response_url")
        status_url = payload.get("status_url")
        if not response_url or not status_url:
            return None

        for _ in range(12):
            with urllib.request.urlopen(
                urllib.request.Request(status_url, headers={"Authorization": f"Key {fal_key}"}),
                timeout=20,
            ) as response:
                status_payload = json.loads(response.read().decode("utf-8"))
            status = str(status_payload.get("status", "")).upper()
            if status == "COMPLETED":
                break
            if status == "FAILED":
                return None
            time.sleep(1)
        else:
            return None

        with urllib.request.urlopen(
            urllib.request.Request(response_url, headers={"Authorization": f"Key {fal_key}"}),
            timeout=20,
        ) as response:
            result_payload = json.loads(response.read().decode("utf-8"))

        image_url = (
            (result_payload.get("images") or [{}])[0].get("url")
            or (result_payload.get("output") or [{}])[0].get("url")
        )
        if not image_url:
            return None

        with urllib.request.urlopen(image_url, timeout=30) as response:
            image_path.write_bytes(response.read())
        return VisualAsset(path=image_path, source="fal", description=prompt)
    except Exception:
        return None


def build_illustration_prompt(slide: dict) -> str:
    title = str(slide.get("title", "")).strip()
    blocks = [block["text"] for block in parse_markdown_blocks(str(slide.get("content_md", ""))) if block["text"]]
    summary = "; ".join(blocks[:3])
    visual_note = str(slide.get("visual", "")).strip()
    layout = normalize_layout(slide.get("layout", ""))
    return (
        "Create a clean presentation illustration for a PowerPoint slide. "
        f"Title: {title}. "
        f"Layout type: {layout}. "
        f"Key points: {summary}. "
        f"Visual direction: {visual_note}. "
        "Use a simple professional style, readable composition, and no embedded text."
    )


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "slide"
